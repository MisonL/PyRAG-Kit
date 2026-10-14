import json
import os
from collections.abc import Mapping
from datetime import datetime
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.worksheet.worksheet import Worksheet  # type: ignore[import-untyped]

from ..ui.display_utils import document_content, document_score, document_source
from ..utils.config import get_settings, is_double_convertible, safe_repr
from ..utils.log_manager import get_module_logger

logger = get_module_logger(__name__)


# xlsx 是 XML，**两处**各自拒绝不同的码位，只处理其中一处会让日志文件写出来
# 却打不开：
#
# 1. ``openpyxl`` 的 ``check_string()`` 在 ``append()`` 期拒绝 C0 控制符
#    （``\x09``/``\x0a``/``\x0d`` 是合法空白，不在其中），抛
#    ``IllegalCharacterError``。
# 2. XML 1.0 不允许的码位在 ``save()`` 期被 ``xml.etree`` 写成非法字符引用：
#    **写入成功、读回失败**。``load_workbook`` 抛 ``ParseError``，而调用方
#    ``log_results`` 的 ``except Exception`` 把它吞成一行日志——用户看到「召回
#    测试完成」，日志文件却打不开。
#
# 第 2 类漏掉的正是孤立代理项（``U+D800``–``U+DFFF``）与 ``U+FFFE``／
# ``U+FFFF``。前者经 legacy pickle 路径可达（pickle 是二进制格式，不受 UTF-8
# 约束，实测 ``pickle.loads(pickle.dumps('caf\udce9'))`` 正常往返）；后者是
# **合法 UTF-8 码位**，``read_text(encoding="utf-8")`` 正常接受，``BasicCleaner``
# 的 ``re.sub(r"\s+", " ")`` 也放行。
#
# 替换而不是删除：日志的用途是记录「检索到了什么」，悄悄吃掉字符会让内容与
# 知识库原文不一致而无人察觉。
_ILLEGAL_XLSX_CHARS = str.maketrans(
    {
        code: "\ufffd"
        for code in (
            *range(0x09),
            *range(0x0B, 0x0D),
            *range(0x0E, 0x20),
            *range(0xD800, 0xE000),
            0xFFFE,
            0xFFFF,
        )
    }
)


# 同一秒内的文件名后缀尝试次数。100 个同名日志已经远超真实使用（每次召回测试
# 只构造一个 logger），到这一步说明目录里堆满了同一秒的产物，报错比继续试更合适。
_MAX_LOG_NAME_ATTEMPTS = 100


class ExcelLogger:
    """
    一个用于将召回测试结果记录到 Excel 文件的日志记录器。
    """

    def __init__(self, log_dir: str | None = None):
        """
        初始化 ExcelLogger。

        Args:
            log_dir (str | None): 存储日志文件的目录。省略时使用
                ``Settings.log_path``，与文本日志落在同一处；否则两者会
                因相对路径基于 CWD、绝对路径基于 ROOT_DIR 而分叉。
        """
        if log_dir is None:
            log_dir = get_settings().log_path
        # 确保日志目录存在
        os.makedirs(log_dir, exist_ok=True)

        timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S")
        self.filepath = self._reserve_log_path(log_dir, timestamp)
        try:
            # 创建一个新的 Excel 工作簿和工作表
            self.workbook = Workbook()
            self.worksheet: Worksheet = self.workbook.active  # type: ignore
            self.worksheet.title = "Recall Test Log"

            # 写入表头
            self._write_header()
        except BaseException:
            # 占位文件已经建出来了，构造失败就必须收回去——否则日志目录里会
            # 留下一个 0 字节的 .xlsx，看起来像一次「没有结果的召回测试」。
            try:
                os.unlink(self.filepath)
            except OSError:
                logger.debug("回收占位日志文件失败: %s", self.filepath)
            raise

        logger.info(f"Excel 日志记录器初始化成功，日志将保存至: {self.filepath}")

    @staticmethod
    def _reserve_log_path(log_dir: str, timestamp: str) -> str:
        """独占创建一个带时间戳的日志路径并返回它。

        文件名原先只有到**秒**的时间戳，于是同一秒内构造的第二个实例会拿到
        同一个路径并覆盖前一个文件——实测两个实例的 ``filepath`` 完全相同。
        召回测试的**唯一产出**就是这份日志，覆盖等于悄悄丢掉一次测试的结果。
        可达路径：从召回测试菜单返回主菜单后再次进入（``main.py`` 的菜单循环），
        或嵌入方在同一秒内构造两个 logger。

        判据用 ``O_CREAT | O_EXCL`` 而不是 ``os.path.exists``：后者是「先查后写」，
        两个进程同时构造时会双双认为路径可用、然后互相覆盖。``O_EXCL`` 由内核
        仲裁，冲突时抛 ``FileExistsError``，换下一个后缀重试即可。

        冲突时加 ``_1``、``_2``…… 后缀（第一个仍是原来的名字，保持既有文件名
        形态不变）。
        """
        for suffix in range(_MAX_LOG_NAME_ATTEMPTS):
            name = (
                f"recall_test_log_{timestamp}.xlsx"
                if suffix == 0
                else f"recall_test_log_{timestamp}_{suffix}.xlsx"
            )
            candidate = os.path.join(log_dir, name)
            try:
                handle = os.open(candidate, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            except FileExistsError:
                continue
            os.close(handle)
            return candidate
        raise RuntimeError(
            f"同一秒内已有 {_MAX_LOG_NAME_ATTEMPTS} 个召回测试日志，无法再创建新文件"
            f"（目录: {log_dir}）。请稍后重试。"
        )

    def _write_header(self):
        """写入 Excel 文件的表头。"""
        headers = [
            "时间戳 (Timestamp)",
            "查询 (Query)",
            "排名 (Rank)",
            "综合得分 (Score)",
            "文档来源 (Source)",
            "页码 (Page)",
            "文档内容 (Content)",
            "详细分数 (All Scores)",
        ]
        self.worksheet.append(headers)
        self.workbook.save(self.filepath)

    @staticmethod
    def _metadata_of(document: Any) -> dict[str, Any]:
        """取出条目的 metadata，非映射时返回空 dict。"""
        if not isinstance(document, Mapping):
            return {}
        metadata = document.get("metadata")
        return dict(metadata) if isinstance(metadata, Mapping) else {}

    @staticmethod
    def _raw_field(document: Any, key: str) -> Any:
        """从条目里取一个顶层字段，条目不是映射时返回 ``None``。"""
        if not isinstance(document, Mapping):
            return None
        return document.get(key)

    @staticmethod
    def _cell_value(value: Any) -> Any:
        """把单元格取值收敛成 ``openpyxl`` 能写入的形态。

        ``openpyxl`` 只接受字符串与数字；映射、列表、集合等会抛
        ``ValueError: Cannot convert ... to Excel``。这类值没有可读的表格表示，
        降级成字符串（保留内容以便排查）比抛异常更合适。

        ``int`` 同样不能无条件放行：``openpyxl`` 在写数值前会走
        ``safe_string()``，那里对值调 ``isnan(value)``——超出双精度范围的整数
        （``10**400`` 即 401 位）抛 ``OverflowError: int too large to convert
        to float``。这条比 ``json.dumps`` 的位数上限**更早**触发（双精度上界约
        1.8e308，即 309 位），因此判据用 ``is_double_convertible`` 而不是位数。

        后果比抛异常更隐蔽：``log_results`` 的 ``save`` 外面是
        ``except Exception: logger.exception(...)``，于是整个工作簿**一个字都
        没落盘**，用户只看到日志里一行「保存 Excel 日志文件失败」，召回测试的
        记录全部丢失。在取值处降级，保住其余字段。

        ``str`` 也要过一遍：``openpyxl`` 有**两条独立的拒绝路径**，上面那条是
        ``save()`` 期的数值溢出，另一条是 ``append()`` 期的
        ``IllegalCharacterError``——``check_string()`` 拒绝 ``\x00-\x08``、
        ``\x0b-\x0c``、``\x0e-\x1f`` 这些控制字符。而 ``_append_text_row`` 的
        ``worksheet.append(row)`` 在 ``log_results`` 的 ``try`` **之外**，异常
        直接冒到调用方：一次**成功**的召回被显示成「召回测试出错」，且该批
        **一条都不落盘**（实测 3 条结果落盘 0 条）。

        触发输入完全合法：``\x1b`` 是终端 ANSI 转义，用户把带颜色的
        ``systemctl status``／``git log --color`` 输出粘进 markdown 就是。
        实测 23 个控制码位能穿过 ``BasicCleaner``（``re.sub(r"\\s+", " ")``
        只吃空白类）并让 openpyxl 抛错。

        处理方式与 ``=`` 公式那支一致：**替换**而不是丢弃内容——控制字符没有
        可读的表格表示，替换成可见的 ``\ufffd`` 既保住了「这里原本有东西」，
        也让日志能落盘。用 ``str.translate`` 一次扫完，不逐字符判断。
        """
        if value is None or isinstance(value, float):
            return value
        if isinstance(value, bool):
            return value
        if isinstance(value, str):
            return value.translate(_ILLEGAL_XLSX_CHARS)
        if isinstance(value, int):
            # 降级用 ``safe_repr`` 而不是 ``str``：位数超限的整数连 ``str()``
            # 都会抛（``str(10**4301)`` -> ``ValueError``），用它等于把
            # ``OverflowError`` 换成同一类错误的另一种形态。
            return value if is_double_convertible(value) else safe_repr(value)
        return safe_repr(value).translate(_ILLEGAL_XLSX_CHARS)

    @staticmethod
    def _json_safe(value: Any) -> Any:
        """把要写进 ``json.dumps`` 的分数收敛成可序列化形态。

        ``int`` 不能无条件放行：``json.dumps(10**4301)`` 抛
        ``ValueError: Exceeds the limit (4300 digits)``（CPython 的
        ``sys.get_int_max_str_digits()``，默认 4300），而本方法的职责正是「收敛成
        可序列化形态」。这类值来自第三方 store 的 ``semantic_score``／
        ``keyword_score``——本函数就是为「第三方 store 不经 FaissStore 归一化」
        而存在的，因此必须自己兜住，不能指望上游。

        位数在上限内的整数照常保留原类型（``json.dumps`` 能处理，且日志的
        用途是记录真实取值）；超出上限的降级成字符串，与其它无法序列化的形态
        一致。判据用 ``is_double_convertible`` 不合适——它判的是双精度范围
        （约 1.8e308，即 309 位），比 ``json.dumps`` 的实际容忍度高得多，会把
        ``10**400`` 这种能正常序列化的值也降级掉。
        """
        if value is None or isinstance(value, (str, float, bool)):
            return value
        if isinstance(value, int):
            # 同 ``_cell_value``：位数超限时连 ``str()`` 都会抛，判据走
            # ``safe_repr`` 的实现（超限时用 ``bit_length()`` 给量级）。
            try:
                str(value)
            except ValueError:
                return safe_repr(value)
            return value
        return safe_repr(value)

    def _append_text_row(self, row: list[Any]) -> None:
        """追加一行，并把其中的字符串单元格固定为文本类型。

        ``openpyxl`` 对以 ``=`` 开头的字符串存成**公式**（``data_type='f'``），
        Excel 打开日志时求值。文档内容、来源与查询都来自知识库与用户输入，实测
        ``=1+1`` 落盘后是公式、``=cmd|'/c calc'!A0`` 这类 DDE 载荷同样照单全收。
        两条后果：日志不再忠实记录「检索到了什么」（内容列显示的是计算结果），
        以及在打开日志的机器上执行文档里的内容。

        这里写 ``data_type = "s"`` 而不是给值加前缀（``'=1+1``）：前缀会改写内容，
        而日志的用途正是逐字记录。非字符串（分数、排名）保持原类型不动。

        **本方法是整行的唯一收口点**，因此控制字符也在这里统一过滤，而不是指望
        每个调用方记得处理：``openpyxl`` 的 ``check_string()`` 拒绝
        ``\x00-\x08``／``\x0b-\x0c``／``\x0e-\x1f``，命中即抛
        ``IllegalCharacterError``，而本方法的调用点位于 ``log_results`` 的
        ``try`` **之外**——异常直接冒到调用方，一次**成功**的召回被显示成
        「召回测试出错」，且该批**一条都不落盘**（实测 3 条结果落盘 0 条）。

        触发输入完全合法：``\x1b`` 是终端 ANSI 转义，用户把带颜色的
        ``systemctl status``／``git log --color`` 输出粘进 markdown 就是。
        实测 23 个控制码位能穿过 ``BasicCleaner``（``re.sub(r"\\s+", " ")``
        只吃空白类）。逐列处理会漏掉 ``query`` 这类用户输入列，收口在这里则
        连未来新增的列也一并覆盖。
        """
        sanitized = [
            value.translate(_ILLEGAL_XLSX_CHARS) if isinstance(value, str) else value
            for value in row
        ]
        self.worksheet.append(sanitized)
        last_row = self.worksheet.max_row
        for column, value in enumerate(sanitized, start=1):
            if isinstance(value, str):
                self.worksheet.cell(row=last_row, column=column).data_type = "s"

    def log_results(self, query: str, results: list[dict[str, Any]]):
        """
        将单次查询的结果记录到 Excel 文件中。

        Args:
            query (str):用户的查询。
            results (List[Dict[str, Any]]): 从检索器返回的文档列表。
        """
        timestamp = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S")

        # ``query`` 也要过 ``_cell_value``。它是本方法里唯一没走收敛的列：
        # 其余列分别经 ``document_*``／``_cell_value``／``_json_safe``，而
        # ``query`` 直接进 ``row``。本方法签名声明 ``query: str``，但它是**公开
        # 方法**，且这里是「日志是旁路、失败不该丢记录」这条契约的入口。
        #
        # 实测 ``query=10**4301``（位数超限的整数）时 ``openpyxl`` 在
        # ``save()`` 期对每个数值调 ``safe_string`` -> ``isnan(value)``，抛
        # ``OverflowError: int too large to convert to float``。后果正是
        # ``_cell_value`` 的 docstring 描述的那一种：``save`` 外的
        # ``except Exception`` 把它吞成一行日志，整个工作簿**一个字都没落盘**
        # ——同一次查询里 3 条完全正常的结果一并丢失，用户看到「召回测试完成」
        # 而日志文件是坏的（``load_workbook`` 连 ``[Content_Types].xml`` 都
        # 找不到）。在取值处降级成 ``safe_repr`` 的量级描述，保住整批记录。
        query = self._cell_value(query)

        # ``results`` 的形态同样要判。本方法签名声明 ``list[dict]``，但它是**公开
        # 方法**，而两种非列表取值都不在契约内：
        #   * ``"abc"`` 是合法可迭代对象，下面的 ``enumerate`` 逐**字符**迭代，
        #     每个字符经 ``document_*`` 助手收敛成占位符，于是写出 3 条内容与来源
        #     都是占位符的伪结果行——用户看到日志写着 3 条命中，而实际返回的是
        #     字符串；
        #   * ``5``／``True``／``1.5`` 抛裸 ``TypeError: 'int' object is not
        #     iterable``，而本方法的调用点位于 ``try`` **之外**（见
        #     ``_append_text_row`` 的说明），异常直接冒到召回测试主循环的
        #     ``except Exception``——一次**成功**的召回被显示成「召回测试出错」。
        # 按「非列表一律视为无结果」收敛：写一行「无结果」比写三条占位垃圾诚实，
        # 与 ``if not results`` 的既有分支同一条出口。
        if not isinstance(results, list):
            results = []

        row: list[Any]
        if not results:
            # 如果没有结果，也记录一条信息
            row = [timestamp, query, "N/A", "N/A", "无结果", "N/A", "N/A", "N/A"]
            self._append_text_row(row)
        else:
            for i, doc in enumerate(results):
                rank = i + 1
                # 字段形态统一由 display_utils 的助手收敛，与展示层同一套判据：
                # 检索结果可以来自任何 ``VectorStoreBase`` 实现，第三方 store 不
                # 经 ``FaissStore`` 的载入期归一化。日志是**旁路**——它失败不该让
                # 一次成功的召回测试以异常收场，也不该丢掉其余正常条目的记录。
                #
                # 直接取值会抛异常的地方实测有六处：``metadata`` 为 None／列表
                # 与条目本身不是映射都会 ``AttributeError``；``page`` 是映射／
                # 列表时 ``openpyxl`` 抛 ``ValueError: Cannot convert ... to
                # Excel``；``semantic_score`` 是 set 时 ``json.dumps`` 抛
                # ``TypeError``；``score`` 是映射／列表同样进不了 Excel。
                score = document_score(doc, 0.0)
                content = document_content(doc, "")
                source = document_source(doc, "未知")
                page = self._cell_value(self._metadata_of(doc).get("page", "N/A"))

                # 提取所有可用的分数。``semantic_score``／``keyword_score`` 同样
                # 收敛——它们要经 ``json.dumps``，集合、映射这类形态无法序列化。
                # 条目本身可能不是映射（第三方 store），因此不直接 ``doc.get``。
                all_scores = {
                    "score": score,
                    "semantic_score": self._json_safe(self._raw_field(doc, "semantic_score")),
                    "keyword_score": self._json_safe(self._raw_field(doc, "keyword_score")),
                }
                # 过滤掉值为 None 的分数
                all_scores_str = json.dumps({k: v for k, v in all_scores.items() if v is not None})

                row = [timestamp, query, rank, score, source, page, content, all_scores_str]
                self._append_text_row(row)

        try:
            self.workbook.save(self.filepath)
            logger.debug(f"成功将查询 '{query}' 的 {len(results)} 条结果记录到 {self.filepath}")
        except Exception:
            logger.exception("保存 Excel 日志文件失败")
